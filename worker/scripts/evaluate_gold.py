"""人工核对集评估（ADR-0021 §人工核对）—— 三套系统对 400 条人工标签的准确率，写 `meta_kv.ai_validation`。

    cd worker && python -m scripts.evaluate_gold --file gold-400.xlsx
    cd worker && python -m scripts.evaluate_gold --file gold-llm-400.xlsx                # 只核对 Luna 的那份表
    cd worker && python -m scripts.evaluate_gold --file D:/x/gold-400.xlsx --labels-file D:/x/gold-400-model-labels.xlsx
    cd worker && python -m scripts.evaluate_gold --file gold-400.xlsx --from-db          # 不用标签表，按正文回库取现行结论
    cd worker && python -m scripts.evaluate_gold --file gold-400.xlsx --no-write         # 只算不写库

## 算什么

按编号把人工表与模型标签表合并，对三套系统各算：

- `relevance_accuracy`：三类相关性准确率（人工填了相关性、且该系统判过的行）；
- `attitude_accuracy`：人工判**相关**且填了态度的行里，系统态度与人一致的比例。系统判成无关
  （没有态度）算错 —— 页面上那条评论就是没态度；
- `attitude_macro_f1`：态度三类的 macro-F1（同一批行）；
- 混淆矩阵。

三套：`student`（学生标签）、`llm`（Luna 现行标签）、`combined`（生产上的路由：学生两头概率
都 ≥ 路由阈值用学生，否则用 Luna；Luna 没标过就还是学生；学生没标过就是 Luna）。
`meta_kv.ai_validation` 顶层三个数取 `combined` —— 那是页面上实际展示的那套。

**每套系统只在它判过的行上算**，分母 `n_relevance` 与覆盖率 `coverage` 一起写进报告：Luna 没见过
学生高置信放行的那些评论，把它们记成 Luna 的错是在量一个它没做过的事；反过来 `combined`
判不出的行（两套都没标）照记为错，因为页面上那条评论确实没有结论。

## 模型标签从哪来

1. 默认：同目录 `<表名>-model-labels.xlsx`（`gold_sample.py` 一起写出的那份）。表头**容缺**：只核对
   Luna 的表（`gold_sample --llm-only`）没有学生列，读出来学生一律 None。
2. 找不到标签表、或标签表里一行 Luna／学生标签都没有、或显式 `--from-db`：按 `(产品代码, 评论正文)`
   回库找判定单元，帖子标题与父评论用来在复读评论之间消歧，然后读库里的现行结论（Luna ＝ provider
   不是 rule／local_model／propagated 的链末行；学生 ＝ provider=local_model 的最新一行）。匹配不到或
   仍有歧义的行记进 `n_unmatched`，不猜。

## 形状（另一位工程师在 backend/core/meta.py 读，**不能改**）

    {"level": "spot_check", "n": 400, "date": "YYYY-MM-DD",
     "relevance_accuracy": 0.91, "attitude_accuracy": 0.87, "attitude_macro_f1": 0.85,
     "by_system": {"student": {同三键}, "llm": {同三键}, "combined": {同三键}}}

算不出来的量写 `null`（例如没有人填态度、或整套系统一条都没判过），不写 0（铁律 2）。写完
`bump_revision("annotation")` 让后端缓存失效。这是量尺不是门槛：数字低不会让任何结论下线（ADR-0019）。
"""

import argparse
import json
import logging
import os
import sys
from collections import Counter, defaultdict
from pathlib import Path

from sqlalchemy import insert, select, update

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
REPO_ROOT = Path(__file__).resolve().parents[2]
for cand in (REPO_ROOT, Path(__file__).resolve().parents[1]):
    if (cand / "radar_db").is_dir() and str(cand) not in sys.path:
        sys.path.insert(0, str(cand))

import clock  # noqa: E402
from models import registry  # noqa: E402
from radar_db import default_data_dir, make_engine  # noqa: E402
from radar_db.schema import annotation_runs, annotations, comments, feeds, meta_kv  # noqa: E402

log = logging.getLogger("worker.evaluate_gold")

OUT_DIR = REPO_ROOT / ".scratch" / "llm-90d"
REL_FROM_ZH = {"相关": "relevant", "无关": "irrelevant", "需上下文": "needs_context"}
ATT_FROM_ZH = {"积极": "positive", "消极": "negative", "中性": "neutral"}
REL_LABELS = ("relevant", "irrelevant", "needs_context")
ATT_LABELS = ("positive", "neutral", "negative")
SYSTEMS = ("student", "llm", "combined")
# 不算 Luna 结论的写入方（ADR-0021 §8）：规则行是词表判的，学生行与抄来的行不是 Luna 说的。
NON_LLM_PROVIDERS = ("rule", "local_model", "propagated")

# 模型标签表的列：`编号` 必有，其余容缺（只核对 Luna 的表没有学生列）。
LABEL_COLUMNS = {
    "comment_id": "comment_id", "code": "产品代码", "stratum": "层",
    "student_relevance": "学生相关性", "student_relevance_p": "学生相关性概率",
    "student_attitude": "学生态度", "student_attitude_p": "学生态度概率",
    "llm_relevance": "Luna相关性", "llm_attitude": "Luna态度",
}


def _norm(v, table):
    if v is None:
        return None
    s = str(v).strip()
    if not s:
        return None
    return table.get(s, s if s in table.values() else None)


def _cell(row, idx, name):
    i = idx.get(name)
    if i is None or i >= len(row):
        return None
    v = row[i]
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def read_gold(path):
    """`标注` 页 → `{编号: {relevance, attitude, note, code, text, title, parent}}`。没填相关性的行不进结果。

    后四个键是回库匹配用的原文与上下文（`gold_sample` 写表时的原样），评估本身只看前两个。
    """
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["标注"] if "标注" in wb.sheetnames else wb.active
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows)]
    idx = {h: i for i, h in enumerate(header)}
    for required in ("编号", "相关性"):
        if required not in idx:
            raise ValueError(f"人工表缺少「{required}」列：{list(header)}")
    out = {}
    for r in rows:
        if not r or _cell(r, idx, "编号") is None:
            continue
        rel = _norm(_cell(r, idx, "相关性"), REL_FROM_ZH)
        att = _norm(_cell(r, idx, "态度"), ATT_FROM_ZH)
        if rel is None:
            continue
        code = _cell(r, idx, "产品代码")
        out[str(_cell(r, idx, "编号"))] = {
            "relevance": rel, "attitude": att if rel == "relevant" else None,
            "note": _cell(r, idx, "备注"),
            "code": str(code) if code is not None else None,
            "text": _cell(r, idx, "评论正文"), "title": _cell(r, idx, "帖子标题"), "parent": _cell(r, idx, "父评论"),
        }
    return out


def read_model_labels(path):
    """模型标签表 → `{编号: {comment_id, code, stratum, student: {...}, llm: {...}}}`。

    表头容缺：没有学生列（只核对 Luna 的表）时学生四项全为 None。连 `Luna相关性` 与 `学生相关性`
    都没有的表不是标签表，抛 `ValueError` 让调用方回库。
    """
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows)]
    idx = {h: i for i, h in enumerate(header)}
    if "编号" not in idx:
        raise ValueError(f"模型标签表缺少「编号」列：{list(header)}")
    if LABEL_COLUMNS["llm_relevance"] not in idx and LABEL_COLUMNS["student_relevance"] not in idx:
        raise ValueError(f"模型标签表既没有「Luna相关性」也没有「学生相关性」列：{list(header)}")
    out = {}
    for r in rows:
        if not r or _cell(r, idx, "编号") is None:
            continue
        g = {k: _cell(r, idx, col) for k, col in LABEL_COLUMNS.items()}
        code = g["code"]
        out[str(_cell(r, idx, "编号"))] = {
            "comment_id": g["comment_id"], "code": str(code) if code is not None else None, "stratum": g["stratum"],
            "student": {"relevance": _norm(g["student_relevance"], REL_FROM_ZH),
                        "attitude": _norm(g["student_attitude"], ATT_FROM_ZH),
                        "relevance_p": g["student_relevance_p"], "attitude_p": g["student_attitude_p"]},
            "llm": {"relevance": _norm(g["llm_relevance"], REL_FROM_ZH),
                    "attitude": _norm(g["llm_attitude"], ATT_FROM_ZH)},
        }
    return out


def has_any_label(model, ids=None):
    """标签表里（这些编号）有没有至少一行 Luna 或学生的相关性。全空的表等于没有表。"""
    keys = model.keys() if ids is None else (k for k in ids if k in model)
    return any(model[k]["student"].get("relevance") is not None or model[k]["llm"].get("relevance") is not None
               for k in keys)


# ── 回库取标签 ────────────────────────────────────────────────────────────


def _chunks(items, n=900):
    items = list(items)
    for i in range(0, len(items), n):
        yield items[i:i + n]


def _student_rows(engine, kind, ids):
    """`{(comment_id, code): (value, calibrated_confidence)}`：provider=local_model 的最新一行。"""
    out = {}
    with engine.connect() as conn:
        for chunk in _chunks(ids):
            rows = conn.execute(
                select(annotations.c.target_id, annotations.c.subject_code, annotations.c.value_json,
                       annotations.c.calibrated_confidence)
                .select_from(annotations.join(annotation_runs, annotation_runs.c.run_id == annotations.c.run_id))
                .where(annotations.c.kind == kind, annotations.c.target_type == "comment",
                       annotations.c.target_id.in_(chunk), annotation_runs.c.provider == "local_model",
                       annotations.c.review_state != "rejected")
                .order_by(annotations.c.annotation_id)
            )
            for tid, code, vj, conf in rows:
                out[(tid, code)] = (json.loads(vj), conf)
    return out


def labels_from_db(engine, gold):
    """按 `(产品代码, 评论正文)` 回库找判定单元，读现行结论。返回 `(model, stats)`。

    复读评论（同产品同正文多条）用帖子标题＋父评论消歧；仍分不开且各候选的结论不一致 ⇒ 记
    `ambiguous`，这一行不进评估。匹配不到 ⇒ `unmatched`。两种都不猜。
    """
    from radar_db.annotations_read import current_annotations

    wanted = {k: g for k, g in gold.items() if g.get("code") and g.get("text")}
    texts = sorted({g["text"] for g in wanted.values()})
    parent = comments.alias("parent")
    candidates = defaultdict(list)  # text -> [(comment_id, title, parent_text)]
    with engine.connect() as conn:
        for chunk in _chunks(texts, 400):
            rows = conn.execute(
                select(comments.c.comment_id, comments.c.content, feeds.c.title, parent.c.content.label("parent_text"))
                .select_from(comments.outerjoin(feeds, feeds.c.feed_id == comments.c.feed_id)
                             .outerjoin(parent, parent.c.comment_id == comments.c.reply_to_comment_id))
                .where(comments.c.content.in_(chunk))
            )
            for cid, content, title, parent_text in rows:
                candidates[content].append((cid, title, parent_text))
        provider_of = dict(conn.execute(select(annotation_runs.c.run_id, annotation_runs.c.provider)).all())

    all_ids = sorted({cid for lst in candidates.values() for cid, _t, _p in lst})
    cur_rel = current_annotations(engine, "relevance", "comment", ids=all_ids)
    cur_att = current_annotations(engine, "attitude", "comment", ids=all_ids)
    st_rel = _student_rows(engine, "relevance", all_ids)
    st_att = _student_rows(engine, "attitude", all_ids)

    def luna(cur, unit):
        r = cur.get(unit)
        if r is None or provider_of.get(r["run_id"], "llm") in NON_LLM_PROVIDERS:
            return None
        return r["value"]

    def labels_of(unit):
        rel, rel_p = st_rel.get(unit, (None, None))
        att, att_p = st_att.get(unit, (None, None))
        l_rel = luna(cur_rel, unit)
        return {
            "student": {"relevance": rel, "attitude": att if rel == "relevant" else None,
                        "relevance_p": rel_p, "attitude_p": att_p if rel == "relevant" else None},
            "llm": {"relevance": l_rel, "attitude": luna(cur_att, unit) if l_rel == "relevant" else None},
        }

    def _same(a, b):
        return (a or "").strip() == (b or "").strip()

    model, stats = {}, Counter()
    for gid, g in wanted.items():
        code = g["code"]
        pool = [c for c in candidates.get(g["text"], []) if (c[0], code) in cur_rel or (c[0], code) in st_rel]
        if not pool:
            stats["unmatched"] += 1
            continue
        if len(pool) > 1:
            narrowed = [c for c in pool if _same(c[1], g.get("title")) and _same(c[2], g.get("parent"))]
            pool = narrowed or pool
        labeled = {cid: labels_of((cid, code)) for cid, _t, _p in pool}
        distinct = {json.dumps(v, sort_keys=True) for v in labeled.values()}
        if len(distinct) > 1:
            stats["ambiguous"] += 1
            continue
        cid = pool[0][0]
        model[gid] = {"comment_id": cid, "code": code, "stratum": None, **labeled[cid]}
        stats["matched"] += 1
    stats["skipped_no_text"] = len(gold) - len(wanted)
    return model, dict(stats)


# ── 评估 ──────────────────────────────────────────────────────────────────


def _label_or_none(sysd):
    rel = sysd.get("relevance")
    return {"relevance": rel, "attitude": sysd.get("attitude") if rel == "relevant" else None}


def combined_label(m, threshold=None):
    """生产路由：学生两头都够置信用学生，否则用 Luna；只有一边有标签就用那一边。"""
    threshold = registry.route_threshold() if threshold is None else threshold
    st, lu = m["student"], m["llm"]
    if st.get("relevance") is None:
        return _label_or_none(lu)
    if lu.get("relevance") is None:
        return _label_or_none(st)
    rp, ap_ = st.get("relevance_p"), st.get("attitude_p")
    confident = rp is not None and rp >= threshold and st["relevance"] != "needs_context"
    if confident and st["relevance"] == "relevant":
        confident = ap_ is not None and ap_ >= threshold
    return _label_or_none(st if confident else lu)


def _accuracy(pairs):
    pairs = [(g, p) for g, p in pairs if g is not None]
    if not pairs:
        return None, 0
    return round(sum(1 for g, p in pairs if g == p) / len(pairs), 4), len(pairs)


def _confusion(pairs, labels):
    m = {g: {p: 0 for p in labels + ("none",)} for g in labels}
    for g, p in pairs:
        if g in m:
            m[g][p if p in labels else "none"] += 1
    return m


def _macro_f1(pairs, labels):
    pairs = [(g, p) for g, p in pairs if g is not None]
    if not pairs:
        return None
    f1s = []
    for lab in labels:
        tp = sum(1 for g, p in pairs if g == lab and p == lab)
        fp = sum(1 for g, p in pairs if g != lab and p == lab)
        fn = sum(1 for g, p in pairs if g == lab and p != lab)
        if tp + fp + fn == 0:
            continue  # 这一类人工没标、系统也没判：不参与平均
        prec = tp / (tp + fp) if tp + fp else 0.0
        rec = tp / (tp + fn) if tp + fn else 0.0
        f1s.append(2 * prec * rec / (prec + rec) if prec + rec else 0.0)
    return round(sum(f1s) / len(f1s), 4) if f1s else None


def evaluate(gold, model, *, threshold=None):
    """返回 `{n, n_gold, n_unmatched, by_system: {system: {relevance_accuracy, attitude_accuracy,
    attitude_macro_f1, n_relevance, n_attitude, n_unlabeled, coverage, confusion}}}`。

    `student`／`llm` 只在自己判过的行上算（`n_unlabeled` 记没判的行数）；`combined` 是页面上那套，
    两边都没标的行算错 —— 那条评论在页面上就是没结论。整套系统一行都没判过 ⇒ 三个指标全 None。
    """
    joined = [(gold[k], model[k]) for k in sorted(gold) if k in model]
    out = {"n": len(joined), "n_gold": len(gold), "n_unmatched": len(set(gold) - set(model)), "by_system": {}}
    for system in SYSTEMS:
        rel_pairs, att_pairs, unlabeled = [], [], 0
        for g, m in joined:
            pred = combined_label(m, threshold) if system == "combined" else _label_or_none(m[system])
            if pred.get("relevance") is None:
                unlabeled += 1
                if system != "combined":
                    continue
            rel_pairs.append((g["relevance"], pred.get("relevance")))
            if g["relevance"] == "relevant" and g["attitude"] is not None:
                att_pairs.append((g["attitude"], pred.get("attitude") if pred.get("relevance") == "relevant" else None))
        rel_acc, n_rel = _accuracy(rel_pairs)
        att_acc, n_att = _accuracy(att_pairs)
        # 一行都没判过的系统给 None 不给 0：0 是「量了，全错」，这里是「没有东西可量」。
        labeled_any = any(p is not None for _g, p in rel_pairs)
        out["by_system"][system] = {
            "relevance_accuracy": rel_acc if labeled_any else None,
            "attitude_accuracy": att_acc if labeled_any else None,
            "attitude_macro_f1": _macro_f1(att_pairs, ATT_LABELS) if labeled_any else None,
            "n_relevance": n_rel, "n_attitude": n_att, "n_unlabeled": unlabeled,
            "coverage": round(n_rel / len(joined), 4) if joined else None,
            "confusion": {"relevance": _confusion(rel_pairs, REL_LABELS), "attitude": _confusion(att_pairs, ATT_LABELS)},
        }
    return out


def ai_validation_payload(result, date_str):
    """`meta_kv.ai_validation` 的形状（见模块 docstring）。顶层三个数＝combined。"""
    keys = ("relevance_accuracy", "attitude_accuracy", "attitude_macro_f1")
    comb = result["by_system"]["combined"]
    return {
        "level": "spot_check", "n": result["n"], "date": date_str,
        **{k: comb[k] for k in keys},
        "by_system": {s: {k: result["by_system"][s][k] for k in keys} for s in SYSTEMS},
    }


def write_validation(engine, payload):
    from radar_db.revisions import bump_revision

    value = json.dumps(payload, ensure_ascii=False)
    with engine.begin() as conn:
        if not conn.execute(update(meta_kv).where(meta_kv.c.k == "ai_validation").values(v=value)).rowcount:
            conn.execute(insert(meta_kv).values(k="ai_validation", v=value))
        bump_revision(conn, "annotation")


def _resolve(path, default_dir):
    p = Path(path)
    if p.exists():
        return p
    alt = Path(default_dir) / p.name
    if alt.exists():
        return alt
    raise SystemExit(f"找不到 {path}（也不在 {default_dir}）")


def load_model_labels(gold, gold_path, *, labels_file=None, from_db=False, engine_factory=make_engine):
    """定模型标签来源：标签表 → 回库。返回 `(model, source, stats)`，`source ∈ {labels_file, db}`。"""
    labels_path = None
    if labels_file:
        labels_path = Path(labels_file)
        if not labels_path.exists():
            raise SystemExit(f"找不到模型标签表 {labels_path}")
    else:
        default = gold_path.with_name(gold_path.stem + "-model-labels.xlsx")
        labels_path = default if default.exists() else None
    if labels_path is not None and not from_db:
        try:
            model = read_model_labels(labels_path)
        except ValueError as exc:
            log.warning("%s；改为按正文回库取标签", exc)
        else:
            if has_any_label(model, gold):
                return model, "labels_file", {"labels_file": labels_path.name, "rows": len(model)}
            log.warning("标签表 %s 里没有任何 Luna／学生标签；改为按正文回库取标签", labels_path.name)
    elif labels_path is None and not from_db:
        log.info("同目录没有 %s-model-labels.xlsx，按正文回库取标签", gold_path.stem)
    model, stats = labels_from_db(engine_factory(), gold)
    return model, "db", stats


def main(argv=None):
    ap = argparse.ArgumentParser(description="评估人工核对集，写 meta_kv.ai_validation")
    ap.add_argument("--file", required=True, help="人工填好的 gold-400.xlsx／gold-llm-400.xlsx（相对路径先在数据目录找）")
    ap.add_argument("--labels-file", help="模型标签表，默认同目录下 <表名>-model-labels.xlsx；找不到就回库")
    ap.add_argument("--from-db", action="store_true", help="不读标签表，按 (产品代码, 评论正文) 回库取现行结论")
    ap.add_argument("--no-write", action="store_true", help="只算不写 meta_kv")
    ap.add_argument("--threshold", type=float, help="combined 路由阈值，默认 STUDENT_ROUTE_THRESHOLD")
    args = ap.parse_args(argv)
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)-5s %(name)s %(message)s")

    data_dir = default_data_dir()
    gold_path = _resolve(args.file, data_dir)
    gold = read_gold(gold_path)
    if not gold:
        raise SystemExit("人工表里没有填好的行（相关性列为空）")
    model, source, source_stats = load_model_labels(gold, gold_path, labels_file=args.labels_file, from_db=args.from_db)
    if not model:
        raise SystemExit(f"一行模型标签都没取到（来源 {source}：{source_stats}）；库里还没有这些评论的现行结论？")
    result = evaluate(gold, model, threshold=args.threshold)
    if result["n"] == 0:
        raise SystemExit(f"人工表与模型标签按编号合并后一行都对不上（来源 {source}：{source_stats}）")
    stamp = clock.now()
    payload = ai_validation_payload(result, stamp.strftime("%Y-%m-%d"))
    report = {"stamp": stamp.strftime("%Y%m%dT%H%M%S"), "gold_file": gold_path.name,
              "labels_source": source, "labels_stats": source_stats,
              "by_stratum": dict(Counter(str(model[k].get("stratum")) for k in gold if k in model)),
              "gold_distribution": {
                  "relevance": dict(Counter(g["relevance"] for g in gold.values())),
                  "attitude": dict(Counter(g["attitude"] for g in gold.values() if g["attitude"])),
              },
              **result, "ai_validation": payload}
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    out = OUT_DIR / f"gold-eval-{report['stamp']}.json"
    out.write_text(json.dumps(report, ensure_ascii=False, indent=1), encoding="utf-8")
    if not args.no_write:
        write_validation(make_engine(), payload)
        log.info("meta_kv.ai_validation 已写入并 bump annotation revision")
    print(json.dumps(payload, ensure_ascii=False, indent=1))
    for system in SYSTEMS:
        s = result["by_system"][system]
        print(f"{system:9s} 判过 {s['n_relevance']:>4}/{result['n']} 行（覆盖 {s['coverage']}），"
              f"态度样本 {s['n_attitude']}")
    if result["n_unmatched"]:
        print(f"未匹配 {result['n_unmatched']} 行（不进分母）")
    print(f"\n完整报告：{out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
